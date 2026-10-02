from pathlib import Path
from datetime import datetime, timedelta
import random
import math
import shutil
import zipfile

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SEED = 42
random.seed(SEED)
np.random.seed(SEED)

BASE_DIR = Path('/mnt/data/projeto_faturamento_oncologia_portfolio')
RAW_DIR = BASE_DIR / 'data' / 'raw'
SCRIPTS_DIR = BASE_DIR / 'scripts'
RAW_DIR.mkdir(parents=True, exist_ok=True)
SCRIPTS_DIR.mkdir(parents=True, exist_ok=True)

# -------------------------------
# Cadastros 100% ficticios
# -------------------------------
convenios = [
    (101, 'Operadora Alfa'),
    (102, 'Operadora Beta'),
    (103, 'Operadora Gama'),
    (104, 'Operadora Delta'),
    (105, 'Operadora Epsilon'),
    (106, 'Operadora Zeta'),
]
unidades = ['Unidade Litoral', 'Unidade Serra']
setores = ['Oncologia Clinica', 'Quimioterapia', 'Hematologia', 'Imunobiologicos']
faturas = ['CONSULTA', 'QUIMIOTERAPIA', 'IMUNOBIOLOGICO', 'MEDICAMENTO', 'PROCEDIMENTO']

primeiros = ['Ana','Beatriz','Carla','Daniela','Eduardo','Felipe','Gabriela','Helena','Igor','Joana','Lucas','Marina','Nicolas','Patricia','Rafael','Sofia','Thiago','Valeria','William','Yasmin']
sobrenomes = ['Almeida','Barros','Campos','Duarte','Esteves','Ferreira','Gomes','Henrique','Lima','Macedo','Nogueira','Oliveira','Pereira','Queiroz','Ramos','Silva','Teixeira','Vieira']

n_pacientes = 520
pacientes = []
used_names = set()
for pid in range(100001, 100001+n_pacientes):
    while True:
        nome = f"{random.choice(primeiros)} {random.choice(sobrenomes)} {random.choice(sobrenomes)}"
        if nome not in used_names:
            used_names.add(nome)
            break
    pacientes.append((pid, nome))

start = datetime(2025,1,1)
end = datetime(2026,8,31)
num_days = (end-start).days + 1
n = 2200

# Helper para valores realistas por tipo de faturamento
def gerar_valor(tipo):
    if tipo == 'CONSULTA':
        v = np.random.lognormal(mean=math.log(450), sigma=0.25)
    elif tipo == 'PROCEDIMENTO':
        v = np.random.lognormal(mean=math.log(2800), sigma=0.60)
    elif tipo == 'QUIMIOTERAPIA':
        v = np.random.lognormal(mean=math.log(8500), sigma=0.75)
    elif tipo == 'IMUNOBIOLOGICO':
        v = np.random.lognormal(mean=math.log(18000), sigma=0.65)
    else:  # MEDICAMENTO
        v = np.random.lognormal(mean=math.log(14000), sigma=0.80)
    return round(float(min(max(v, 180), 120000)), 2)

rows_fat = []
rows_ctrl = []

for i in range(n):
    id_fat = 500001 + i
    cd_conv, nm_conv = random.choices(convenios, weights=[26,22,18,14,12,8], k=1)[0]
    tipo = random.choices(faturas, weights=[18,34,14,18,16], k=1)[0]
    unidade = random.choices(unidades, weights=[58,42], k=1)[0]
    if tipo == 'CONSULTA':
        setor = 'Oncologia Clinica'
    elif tipo in ('QUIMIOTERAPIA','MEDICAMENTO'):
        setor = random.choices(['Quimioterapia','Oncologia Clinica'], weights=[85,15], k=1)[0]
    elif tipo == 'IMUNOBIOLOGICO':
        setor = 'Imunobiologicos'
    else:
        setor = random.choice(setores)

    dt_at = start + timedelta(days=random.randrange(num_days))
    dt_comp = datetime(dt_at.year, dt_at.month, 1)
    cd_remessa = 80000 + (i // random.choice([8,9,10,11,12]))
    cd_conta = 700000 + i
    cd_atendimento = 1200000 + i
    nr_guia = f"G{2025 if dt_at.year==2025 else 2026}{100000+i}"
    pid, nome = random.choice(pacientes)
    vl_total = gerar_valor(tipo)

    # 87% ja tiveram envio; 13% ainda aguardam envio
    tem_envio = random.random() < 0.87
    dt_envio = dt_at + timedelta(days=random.randint(3, 28)) if tem_envio else pd.NaT

    # Pequenas inconsistencias textuais de proposito para tratamento no Python
    nm_conv_raw = nm_conv
    r = random.random()
    if r < 0.012:
        nm_conv_raw = nm_conv.upper()
    elif r < 0.024:
        nm_conv_raw = nm_conv + ' '

    rows_fat.append({
        'ID_FATURAMENTO': id_fat,
        'CD_CONVENIO': cd_conv,
        'NM_CONVENIO': nm_conv_raw,
        'DS_FATURA': tipo,
        'DT_COMPETENCIA': dt_comp,
        'CD_REMESSA': cd_remessa,
        'DT_ENVIO_SISTEMA': dt_envio,
        'CD_CONTA': cd_conta,
        'CD_ATENDIMENTO': cd_atendimento,
        'NR_GUIA': nr_guia,
        'CD_PACIENTE': pid,
        'NM_PACIENTE': nome,
        'DT_ATENDIMENTO': dt_at,
        'VL_TOTAL_CONTA': vl_total,
        'UNIDADE': unidade,
        'SETOR': setor,
    })

    # Controle manual existe para quase todas as contas; algumas ficam sem lancamento
    if random.random() < 0.95:
        lote_manual = (dt_envio + timedelta(days=random.choice([0,0,0,1,2]))) if tem_envio else pd.NaT
        glosa = 0.0
        data_glosa = pd.NaT
        glosa_desconsiderada = 0.0
        recursado = 0.0
        data_recurso = pd.NaT
        pago_recurso = 0.0
        data_pag_recurso = pd.NaT
        outros = 0.0
        valor_credito = 0.0
        pago_em = pd.NaT
        processo = ''
        taxa_operadora = 0.0
        nf = ''
        aliquota = 0.0
        tributos = 0.0
        status = 'AGUARDANDO ENVIO' if not tem_envio else 'EM ABERTO'

        if tem_envio:
            teve_glosa = random.random() < 0.19
            if teve_glosa:
                glosa = round(vl_total * random.uniform(0.03, 0.28), 2)
                data_glosa = lote_manual + timedelta(days=random.randint(20,70))
                # pequena parte da glosa e desconsiderada administrativamente
                if random.random() < 0.18:
                    glosa_desconsiderada = round(glosa * random.uniform(0.05,0.35), 2)
                # recurso em boa parte das glosas
                if random.random() < 0.72:
                    recursado = round(max(glosa-glosa_desconsiderada,0) * random.uniform(0.65,1.00),2)
                    data_recurso = data_glosa + timedelta(days=random.randint(5,25))
                    if random.random() < 0.70:
                        pago_recurso = round(recursado * random.uniform(0.35,0.95),2)
                        data_pag_recurso = data_recurso + timedelta(days=random.randint(15,55))
                        processo = f"REC-{random.randint(10000,99999)}"

            valor_base_pagavel = max(vl_total - glosa, 0)
            # 72% das remessas ja receberam credito inicial
            if random.random() < 0.72:
                fator_pag = random.uniform(0.96,1.00)
                valor_credito = round(valor_base_pagavel * fator_pag,2)
                pago_em = lote_manual + timedelta(days=random.randint(20,85))
                status = 'PAGO' if glosa == 0 else ('PAGO COM GLOSA' if pago_recurso == 0 else 'PAGO + RECURSO')
            else:
                status = 'EM ABERTO' if glosa == 0 else 'GLOSA EM TRATAMENTO'

            # taxas, tributos e nota apenas em contas com pagamento
            if valor_credito > 0:
                taxa_operadora = round(valor_credito * random.choice([0,0,0,0.005,0.01,0.015]),2)
                aliquota = random.choice([0.0465, 0.06, 0.065])
                tributos = round((valor_credito + pago_recurso) * aliquota,2)
                nf = f"NF{random.randint(100000,999999)}"
                if random.random() < 0.035:
                    outros = round(random.uniform(10,250),2)

        rows_ctrl.append({
            'NR_GUIA': nr_guia,
            'LOTE_ENVIADO_EM': lote_manual,
            'VALOR_CREDITO': valor_credito,
            'PAGO_EM': pago_em,
            'GLOSA_INICIAL': glosa,
            'DATA_GLOSA': data_glosa,
            'GLOSA_DESCONSIDERADA': glosa_desconsiderada,
            'VALOR_RECURSADO': recursado,
            'DATA_RECURSO': data_recurso,
            'VALOR_PAGO_RECURSO': pago_recurso,
            'DATA_PAGAMENTO_RECURSO': data_pag_recurso,
            'OUTRAS_BAIXAS': outros,
            'PROCESSO_PROTOCOLO': processo,
            'TAXA_OPERADORA': taxa_operadora,
            'NUMERO_NF': nf,
            'TRIBUTOS_NF': tributos,
            'ALIQUOTA_TRIBUTOS': aliquota,
            'STATUS_MANUAL': status,
        })

# DataFrames
df_fat = pd.DataFrame(rows_fat)
df_ctrl = pd.DataFrame(rows_ctrl)

# Introduzimos poucos duplicados intencionais na exportacao para tratamento didatico
idx_dup = np.random.choice(df_fat.index, size=8, replace=False)
df_fat = pd.concat([df_fat, df_fat.loc[idx_dup]], ignore_index=True)
# embaralhar
_df_order = np.random.permutation(len(df_fat))
df_fat = df_fat.iloc[_df_order].reset_index(drop=True)

file1 = RAW_DIR / '01_faturamento_sistema.xlsx'
file2 = RAW_DIR / '02_controle_recebimentos_manual.xlsx'

with pd.ExcelWriter(file1, engine='openpyxl') as writer:
    df_fat.to_excel(writer, sheet_name='FATURAMENTO', index=False)
    pd.DataFrame({
        'INFORMACAO': [
            'Base 100% ficticia criada para projeto de portfolio.',
            'Simula uma exportacao do sistema de faturamento de uma clinica oncologica.',
            'Existem duplicidades e pequenas inconsistencias de texto propositalmente para praticar tratamento de dados.',
            'Todos os nomes, codigos, valores e datas foram gerados artificialmente.'
        ]
    }).to_excel(writer, sheet_name='LEIA_ME', index=False)

with pd.ExcelWriter(file2, engine='openpyxl') as writer:
    df_ctrl.to_excel(writer, sheet_name='CONTROLE_MANUAL', index=False)
    pd.DataFrame({
        'INFORMACAO': [
            'Base 100% ficticia criada para projeto de portfolio.',
            'Simula uma planilha preenchida manualmente para acompanhar recebimentos, glosas, recursos, impostos e notas fiscais.',
            'A chave principal para cruzamento com a base de faturamento e NR_GUIA.',
            'Algumas guias nao aparecem nesta planilha propositalmente para praticar validacao de integridade.'
        ]
    }).to_excel(writer, sheet_name='LEIA_ME', index=False)

# Estilizacao simples e profissional das bases brutas
for path, sheet in [(file1,'FATURAMENTO'), (file2,'CONTROLE_MANUAL')]:
    wb = load_workbook(path)
    ws = wb[sheet]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill('solid', fgColor='1F4E78')
    for cell in ws[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    # Table
    ref = ws.dimensions
    tab = Table(displayName='tb_' + ('faturamento' if '01_' in path.name else 'controle'), ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)
    # widths
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in list(col_cells)[:300])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 28)
    # formats
    header_map = {c.value:c.column for c in ws[1]}
    for h,col in header_map.items():
        if h and (h.startswith('DT_') or h.startswith('DATA_') or h in ('PAGO_EM','LOTE_ENVIADO_EM')):
            for cell in ws.iter_cols(min_col=col,max_col=col,min_row=2):
                for x in cell: x.number_format='dd/mm/yyyy'
        if h and (h.startswith('VL_') or 'VALOR' in h or 'GLOSA' in h or 'TRIBUTOS' in h or 'TAXA_OPERADORA'==h or 'OUTRAS_BAIXAS'==h):
            if h != 'ALIQUOTA_TRIBUTOS':
                for cell in ws.iter_cols(min_col=col,max_col=col,min_row=2):
                    for x in cell: x.number_format='R$ #,##0.00;[Red](R$ #,##0.00);-'
        if h == 'ALIQUOTA_TRIBUTOS':
            for cell in ws.iter_cols(min_col=col,max_col=col,min_row=2):
                for x in cell: x.number_format='0.00%'
    ws.sheet_view.showGridLines = False
    # LEIA_ME style
    info = wb['LEIA_ME']
    info.column_dimensions['A'].width = 110
    info['A1'].font = Font(bold=True, color='FFFFFF')
    info['A1'].fill = header_fill
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    info.sheet_view.showGridLines = False
    wb.save(path)

readme = BASE_DIR / 'README_INICIO.md'
readme.write_text('''# Analise Financeira do Faturamento de uma Clinica Oncologica\n\nProjeto de portfolio com dados **100% ficticios**, inspirado em um processo real de acompanhamento de faturamento, recebimentos e glosas.\n\n## Cenario\nA area de faturamento fornece uma exportacao do sistema e mantem uma segunda planilha manual para acompanhar recebimentos, glosas, recursos, impostos e notas fiscais. O objetivo do projeto e tratar e integrar essas fontes com Python e criar uma base analitica para um dashboard financeiro.\n\n## Fontes\n- `data/raw/01_faturamento_sistema.xlsx`: exportacao simulada do sistema.\n- `data/raw/02_controle_recebimentos_manual.xlsx`: controle manual simulado.\n\n## Principais analises planejadas\n- Faturamento bruto e liquido\n- Valor recebido e valor em aberto\n- Glosa inicial, glosa liquida e recuperacao de glosas\n- Prazo medio de envio\n- Prazo medio de recebimento / aging\n- Receita por competencia, operadora, unidade e tipo de faturamento\n- Faixas de atraso (0-30, 31-60, 61-90 e 90+ dias)\n- Indicadores de qualidade do processo: contas sem envio, sem controle manual e duplicidades\n\n> Nenhum dado real de paciente, clinica ou operadora foi utilizado nestes arquivos.\n''', encoding='utf-8')

requirements = BASE_DIR / 'requirements.txt'
requirements.write_text('pandas>=2.2\nopenpyxl>=3.1\nnumpy>=2.0\n', encoding='utf-8')

# copia este proprio gerador para a pasta scripts para transparencia/reprodutibilidade
shutil.copy('/mnt/data/gerar_projeto_faturamento.py', SCRIPTS_DIR / '00_gerar_dados_ficticios.py')

# Zip
zip_path = Path('/mnt/data/projeto_faturamento_oncologia_portfolio_inicial.zip')
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as z:
    for p in BASE_DIR.rglob('*'):
        if p.is_file():
            z.write(p, p.relative_to(BASE_DIR.parent))

print('Criado:', file1)
print('Linhas faturamento:', len(df_fat), 'unicas por ID:', df_fat['ID_FATURAMENTO'].nunique())
print('Criado:', file2)
print('Linhas controle:', len(df_ctrl))
print('Zip:', zip_path)
print('Faturamento total sintetico:', round(df_fat.drop_duplicates('ID_FATURAMENTO')['VL_TOTAL_CONTA'].sum(),2))
print('Periodo:', df_fat['DT_ATENDIMENTO'].min().date(), 'a', df_fat['DT_ATENDIMENTO'].max().date())
